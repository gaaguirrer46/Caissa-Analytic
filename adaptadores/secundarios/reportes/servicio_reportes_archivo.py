"""Adaptador de generación de reportes en PDF, Excel, JSON y HTML.

- Excel: ``openpyxl`` (dependencia de la aplicación, no del dominio).
- PDF: escritor mínimo basado en PDF 1.4 y la fuente Helvetica, sin dependencias
 externas. El contenido se pagina y se codifica en WinAnsi.
- JSON: serie lista para cualquier librería de gráficos.
- HTML: gráfico interactivo autocontenido con
 SVG + JavaScript nativo, sin dependencias externas ni red."""
from __future__ import annotations

import html
import json
import textwrap
from datetime import datetime
from pathlib import Path
from typing import Iterable, List

from dominio.indicadores.entidades import FormatoReporte, Reporte
from dominio.indicadores.puertos import ServicioGeneracionReportes

_ANCHO_LINEA = 90
_LINEAS_POR_PAGINA = 50
_EXTENSIONES = {
    FormatoReporte.PDF: ".pdf",
    FormatoReporte.EXCEL: ".xlsx",
    FormatoReporte.JSON: ".json",
    FormatoReporte.HTML: ".html",
}

def _slug(texto: str) -> str:
    limpio = "".join(
        caracter if caracter.isalnum() or caracter in " -_" else "_"
        for caracter in texto
    ).strip()
    return "_".join(limpio.split()) or "reporte"

def _escapar_pdf(texto: str) -> str:
    return (
        texto.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
        .replace("\r", "")
        .replace("\n", " ")
    )

def _envolver(lineas: Iterable[str]) -> List[str]:
    resultado: List[str] = []
    for linea in lineas:
        partes = textwrap.wrap(str(linea), width=_ANCHO_LINEA) or [""]
        resultado.extend(partes)
    return resultado

def _flujo_contenido(lineas: List[str]) -> bytes:
    comandos = ["BT", "/F1 11 Tf", "14 TL", "50 800 Td"]
    for linea in lineas:
        comandos.append(f"({_escapar_pdf(linea)}) Tj")
        comandos.append("T*")
    comandos.append("ET")
    return "\n".join(comandos).encode("cp1252", "replace")

def _pdf(lineas: List[str]) -> bytes:
    paginas = [
        lineas[i : i + _LINEAS_POR_PAGINA]
        for i in range(0, len(lineas), _LINEAS_POR_PAGINA)
    ] or [[""]]

    ids_contenido = []
    ids_pagina = []
    siguiente = 4
    for _ in paginas:
        ids_contenido.append(siguiente)
        siguiente += 1
        ids_pagina.append(siguiente)
        siguiente += 1

    objetos: List[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (
            f"<< /Type /Pages /Kids "
            f"[{' '.join(f'{pid} 0 R' for pid in ids_pagina)}] "
            f"/Count {len(ids_pagina)} >>"
        ).encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]

    for pagina, id_contenido, id_pagina in zip(paginas, ids_contenido, ids_pagina):
        flujo = _flujo_contenido(pagina)
        objetos.append(
            f"<< /Length {len(flujo)} >>\nstream\n".encode("ascii")
            + flujo
            + b"\nendstream"
        )
        objetos.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {id_contenido} 0 R >>"
            ).encode("ascii")
        )

    salida = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for numero, objeto in enumerate(objetos, start=1):
        offsets.append(len(salida))
        salida += f"{numero} 0 obj\n".encode("ascii")
        salida += objeto
        salida += b"\nendobj\n"

    inicio_xref = len(salida)
    salida += f"xref\n0 {len(objetos) + 1}\n".encode("ascii")
    salida += b"0000000000 65535 f \n"
    for offset in offsets:
        salida += f"{offset:010d} 00000 n \n".encode("ascii")
    salida += (
        f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\n"
        f"startxref\n{inicio_xref}\n%%EOF\n"
    ).encode("ascii")
    return bytes(salida)

_PLANTILLA_GRAFICO = """<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITULO__</title>
<style>
  body { font-family: system-ui, sans-serif; margin: 2rem; color: #1f2937; }
  h1 { font-size: 1.25rem; }
  .meta { color: #6b7280; font-size: .875rem; margin-bottom: 1rem; }
  #chart { position: relative; max-width: 960px; }
  svg { width: 100%; height: auto; background: #fff; border: 1px solid #e5e7eb; border-radius: 8px; }
  .punto { cursor: pointer; }
  #tooltip { position: absolute; pointer-events: none; background: #111827; color: #fff;
    padding: .35rem .6rem; border-radius: 6px; font-size: .8rem; display: none; white-space: nowrap; }
  table { border-collapse: collapse; margin-top: 1.5rem; font-size: .875rem; }
  th, td { border: 1px solid #e5e7eb; padding: .35rem .7rem; text-align: left; }
</style>
</head>
<body>
<h1>__TITULO__</h1>
<p class="meta">Generado: __GENERADO__</p>
<div id="chart"><div id="tooltip"></div></div>
<table id="tabla"></table>
<script>
const DATOS = __DATOS__;
const serie = DATOS.series || [];
const W = 960, H = 420, M = {top: 20, right: 30, bottom: 60, left: 70};
const iw = W - M.left - M.right, ih = H - M.top - M.bottom;
const valores = serie.map(p => Number(p.valor) || 0);
const min = Math.min(...valores), max = Math.max(...valores);
const rango = (max - min) || 1;
const x = i => M.left + (serie.length > 1 ? (i * iw) / (serie.length - 1) : iw / 2);
const y = v => M.top + ih - ((v - min) / rango) * ih;
let svg = `<svg viewBox="0 0 ${W} ${H}" role="img" aria-label="Serie de tendencia">`;
// Eje Y: 5 marcas
for (let k = 0; k <= 4; k++) {
  const valor = min + (rango * k) / 4, py = y(valor);
  svg += `<line x1="${M.left}" y1="${py}" x2="${W - M.right}" y2="${py}" stroke="#f3f4f6"/>`;
  svg += `<text x="${M.left - 8}" y="${py + 4}" text-anchor="end" font-size="12" fill="#6b7280">${valor.toFixed(1)}</text>`;
}
// Línea y puntos
const camino = serie.map((p, i) => `${i ? "L" : "M"}${x(i)},${y(Number(p.valor) || 0)}`).join(" ");
svg += `<path d="${camino}" fill="none" stroke="#2563eb" stroke-width="2.5"/>`;
serie.forEach((p, i) => {
  svg += `<circle class="punto" cx="${x(i)}" cy="${y(Number(p.valor) || 0)}" r="5" fill="#2563eb" data-i="${i}"/>`;
  svg += `<text x="${x(i)}" y="${H - 34}" text-anchor="middle" font-size="11" fill="#6b7280" transform="rotate(-30 ${x(i)} ${H - 34})">${p.etiqueta}</text>`;
});
svg += `<line x1="${M.left}" y1="${M.top + ih}" x2="${W - M.right}" y2="${M.top + ih}" stroke="#9ca3af"/>`;
svg += `</svg>`;
document.getElementById("chart").insertAdjacentHTML("beforeend", svg);
// Tooltip interactivo
const tooltip = document.getElementById("tooltip");
document.querySelectorAll(".punto").forEach(circ => {
  circ.addEventListener("mousemove", ev => {
    const p = serie[Number(circ.dataset.i)];
    tooltip.style.display = "block";
    tooltip.style.left = (ev.offsetX + 14) + "px";
    tooltip.style.top = (ev.offsetY + 14) + "px";
    tooltip.textContent = `${p.etiqueta}: ${p.valor}`;
    circ.setAttribute("r", "7");
  });
  circ.addEventListener("mouseleave", () => {
    tooltip.style.display = "none";
    circ.setAttribute("r", "5");
  });
});
// Tabla de datos (accesibilidad)
const filas = serie.map(p => `<tr><td>${p.etiqueta}</td><td>${p.desde}</td><td>${p.hasta}</td><td>${p.valor}</td></tr>`).join("");
document.getElementById("tabla").innerHTML =
  `<thead><tr><th>Periodo</th><th>Desde</th><th>Hasta</th><th>Valor</th></tr></thead><tbody>${filas}</tbody>`;
</script>
</body>
</html>
"""

class ServicioReportesArchivo(ServicioGeneracionReportes):
    """Genera reportes PDF/Excel/JSON/HTML en disco desde el modelo de dominio."""

    def generar(self, reporte: Reporte, ruta_destino: str) -> str:
        destino = Path(ruta_destino)
        extension = _EXTENSIONES[reporte.formato]
        if destino.suffix.lower() == extension:
            archivo = destino
        else:
            archivo = destino / f"{_slug(reporte.titulo)}{extension}"
        archivo.parent.mkdir(parents=True, exist_ok=True)

        if reporte.formato == FormatoReporte.PDF:
            archivo.write_bytes(_pdf(self._lineas(reporte)))
        elif reporte.formato == FormatoReporte.EXCEL:
            self._excel(reporte, archivo)
        elif reporte.formato == FormatoReporte.JSON:
            self._json(reporte, archivo)
        elif reporte.formato == FormatoReporte.HTML:
            self._html(reporte, archivo)
        else:
            raise ValueError(f"Formato de reporte no soportado: {reporte.formato}")
        return str(archivo)

    @staticmethod
    def _lineas(reporte: Reporte) -> List[str]:
        contenido = reporte.contenido
        columnas = contenido.get("columnas")
        filas = contenido.get("filas")

        es_tabla = isinstance(columnas, list) and isinstance(filas, list)
        lineas = [
            f"{clave}: {valor}"
            for clave, valor in contenido.items()
            if not es_tabla or clave not in ("columnas", "filas")
        ]
        if es_tabla:
            lineas.append(" | ".join(str(columna) for columna in columnas))
            lineas.extend(
                " | ".join(str(valor) for valor in fila) for fila in filas
            )
        return _envolver(
            [
                reporte.titulo,
                "",
                f"Tipo: {reporte.tipo.value}",
                f"Formato: {reporte.formato.value}",
                f"Generado: {reporte.generado_en.isoformat()}",
                "",
                "Secciones:",
                *[f"- {seccion}" for seccion in reporte.secciones],
                "",
                "Contenido:",
                *lineas,
            ]
        )

    @staticmethod
    def _json(reporte: Reporte, archivo: Path) -> None:
        """exporta la serie lista para cualquier librería de gráficos."""
        payload = {
            "titulo": reporte.titulo,
            "tipo": reporte.tipo.value,
            "formato": reporte.formato.value,
            "generado_en": reporte.generado_en.isoformat(),
            "secciones": list(reporte.secciones),
            "contenido": reporte.contenido,
        }
        archivo.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )

    @staticmethod
    def _html(reporte: Reporte, archivo: Path) -> None:
        """gráfico interactivo autocontenido (SVG + JS nativo)."""
        contenido = reporte.contenido
        datos = {
            "titulo": reporte.titulo,
            "indicador": str(contenido.get("indicador", "")),
            "periodicidad": str(contenido.get("periodicidad", "")),
            "desde": str(contenido.get("desde", "")),
            "hasta": str(contenido.get("hasta", "")),
            "series": contenido.get("series", []),
            "resumen": contenido.get("resumen", {}),
        }

        serie_json = json.dumps(datos, ensure_ascii=False, default=str).replace(
            "</", "<\\/"
        )
        titulo = html.escape(reporte.titulo)
        generado = html.escape(reporte.generado_en.isoformat())
        documento = _PLANTILLA_GRAFICO.replace("__TITULO__", titulo).replace(
            "__GENERADO__", generado
        ).replace("__DATOS__", serie_json)
        archivo.write_text(documento, encoding="utf-8")

    @staticmethod
    def _excel(reporte: Reporte, archivo: Path) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Reporte"

        hoja["A1"] = reporte.titulo
        hoja["A1"].font = Font(bold=True, size=14)
        hoja["A2"] = "Tipo"
        hoja["B2"] = reporte.tipo.value
        hoja["A3"] = "Formato"
        hoja["B3"] = reporte.formato.value
        hoja["A4"] = "Generado"
        hoja["B4"] = reporte.generado_en.isoformat()

        fila = 6
        hoja.cell(row=fila, column=1, value="Secciones").font = Font(bold=True)
        fila += 1
        for seccion in reporte.secciones:
            hoja.cell(row=fila, column=1, value=seccion)
            fila += 1

        fila += 1
        hoja.cell(row=fila, column=1, value="Contenido").font = Font(bold=True)
        fila += 1
        columnas = reporte.contenido.get("columnas")
        filas = reporte.contenido.get("filas")

        es_tabla = isinstance(columnas, list) and isinstance(filas, list)
        omitidas = {"columnas", "filas"} if es_tabla else set()
        for clave, valor in reporte.contenido.items():
            if clave in omitidas:
                continue
            hoja.cell(row=fila, column=1, value=str(clave))
            hoja.cell(row=fila, column=2, value=str(valor))
            fila += 1
        if es_tabla:
            fila += 1
            for indice, nombre in enumerate(columnas, start=1):
                hoja.cell(row=fila, column=indice, value=str(nombre)).font = Font(
                    bold=True
                )
            for registro in filas:
                fila += 1
                for indice, valor in enumerate(registro, start=1):
                    hoja.cell(row=fila, column=indice, value=valor)

        hoja.column_dimensions["A"].width = 32
        hoja.column_dimensions["B"].width = 80
        libro.save(archivo)
